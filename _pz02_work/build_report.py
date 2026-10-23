from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from pathlib import Path
import textwrap

root = Path(r'C:\Users\eg0rg\Desktop\quizzy')
work = root / '_pz02_work'
figdir = work / 'figures'
figdir.mkdir(parents=True, exist_ok=True)
template = Path(r'C:\Users\eg0rg\Desktop\кафтан\МДК_01_04_практика_02_шаблон_отчета.docx')
out = root / 'docs' / 'Отчет_ПЗ2_Модульная_архитектура.docx'

font_dir = Path(r'C:\Windows\Fonts')
font_regular = ImageFont.truetype(str(font_dir / 'consola.ttf'), 25)
font_bold = ImageFont.truetype(str(font_dir / 'consolab.ttf'), 29)
font_ui = ImageFont.truetype(str(font_dir / 'arial.ttf'), 27)
font_ui_bold = ImageFont.truetype(str(font_dir / 'arialbd.ttf'), 30)


def code_card(path, title, sections, width=1700):
    pad, title_h, line_h = 44, 72, 37
    lines = []
    for label, source in sections:
        lines.append(('label', label))
        lines.extend(('code', x) for x in source.splitlines())
        lines.append(('gap', ''))
    height = pad * 2 + title_h + len(lines) * line_h
    img = Image.new('RGB', (width, height), '#111827')
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((16, 16, width-16, height-16), radius=20, fill='#111827', outline='#334155', width=2)
    d.text((pad, pad), title, font=font_ui_bold, fill='#e2e8f0')
    y = pad + title_h
    for kind, text in lines:
        if kind == 'label':
            d.text((pad, y), text, font=font_bold, fill='#7dd3fc')
        elif kind == 'code':
            d.text((pad, y), text, font=font_regular, fill='#e5e7eb')
        y += line_h
    img.save(path)

code_card(figdir/'figure1.png', 'Публичный контракт и рабочая реализация', [
    ('QuizMvp/Authentication/IAuthenticationModule.cs', '''public interface IAuthenticationModule
{
    bool RegistrationEnabled { get; }

    Task<RegistrationResult> RegisterAsync(
        RegistrationRequest request,
        CancellationToken cancellationToken = default);

    Task<AuthenticationResult> AuthenticateAsync(
        LoginRequest request,
        CancellationToken cancellationToken = default);
}'''),
    ('QuizMvp/Services/AuthenticationModule.cs', '''if (user is null || string.IsNullOrEmpty(password)
    || !VerifyPassword(password, user.PasswordHash))
{
    return new AuthenticationResult(
        AuthenticationStatus.InvalidCredentials, null);
}
return new AuthenticationResult(
    AuthenticationStatus.Authenticated, ToAccount(user));''')
])

code_card(figdir/'figure2.png', 'Документация контракта и XML-комментарии', [
    ('XML-документация публичного API', '''/// Проверяет имя и пароль. При отсутствии пользователя
/// или несовпадении пароля возвращает InvalidCredentials.
Task<AuthenticationResult> AuthenticateAsync(
    LoginRequest request,
    CancellationToken cancellationToken = default);'''),
    ('docs/authentication-module-contract.md', '''Вход: имя и пароль; допускаются null.
Успех: AuthenticatedAccount (ID, имя, роль).
Неверные данные: общий исход InvalidCredentials.
Побочный эффект: чтения без изменения данных.
Пример вызова и ограничения реализации описаны рядом.''')
])

code_card(figdir/'figure3.png', 'Один потребитель и две реализации на общей границе', [
    ('Потребитель AccountController', '''public AccountController(IAuthenticationModule module)
    => _authenticationModule = module;

await _authenticationModule.AuthenticateAsync(request, token);'''),
    ('Рабочая реализация в Program.cs', '''builder.Services.AddScoped<IAuthenticationModule,
    AuthenticationModule>();'''),
    ('Подмена в тестовой точке сборки', '''services.RemoveAll<IAuthenticationModule>();
services.AddSingleton(authenticationModuleOverride);'''),
    ('Сценарий', 'Один HTTP POST /Account/Login с одинаковыми данными проверяется с обеими реализациями.')
])

# Architecture map illustration
img = Image.new('RGB', (1800, 1060), 'white')
d = ImageDraw.Draw(img)
try:
    map_font = ImageFont.truetype(str(font_dir/'arial.ttf'), 26)
    map_bold = ImageFont.truetype(str(font_dir/'arialbd.ttf'), 27)
except OSError:
    map_font, map_bold = font_ui, font_ui_bold

def box(x, y, w, h, title, subtitle, fill='#eff6ff', outline='#2563eb'):
    d.rounded_rectangle((x,y,x+w,y+h), radius=18, fill=fill, outline=outline, width=3)
    d.text((x+18,y+14), title, font=map_bold, fill='#0f172a')
    yy=y+53
    for line in textwrap.wrap(subtitle, width=max(20,w//15)):
        d.text((x+18,yy), line, font=map_font, fill='#334155'); yy += 34

def arrow(a,b,label=''):
    x1,y1=a; x2,y2=b
    d.line((x1,y1,x2,y2), fill='#475569', width=4)
    import math
    ang=math.atan2(y2-y1,x2-x1); size=16
    p1=(x2-size*math.cos(ang-0.45),y2-size*math.sin(ang-0.45))
    p2=(x2-size*math.cos(ang+0.45),y2-size*math.sin(ang+0.45))
    d.polygon([(x2,y2),p1,p2],fill='#475569')
    if label: d.text(((x1+x2)//2,(y1+y2)//2-32),label,font=map_font,fill='#334155')

box(70,90,500,150,'AccountController','Веб-слой • вызывает контракт')
box(680,90,500,150,'IAuthenticationModule','Публичный контракт • запросы и исходы',fill='#ecfdf5',outline='#059669')
box(1270,90,450,150,'Program.cs','Точка сборки • связывает DI',fill='#fef3c7',outline='#d97706')
box(680,390,500,170,'AuthenticationModule','Рабочая реализация • правила входа/регистрации',fill='#f1f5f9',outline='#475569')
box(90,700,460,160,'ApplicationDbContext','EF Core • хранение пользователей',fill='#f1f5f9',outline='#475569')
box(690,700,450,160,'AuthenticationModuleSettings','Настройка доступности регистрации',fill='#f1f5f9',outline='#475569')
box(1290,700,410,160,'User','Модель записи пользователя',fill='#f1f5f9',outline='#475569')
arrow((570,165),(680,165))
arrow((930,390),(930,240))
arrow((1270,165),(1180,165))
arrow((680,480),(550,700))
arrow((880,560),(880,700))
arrow((1180,480),(1490,700))
arrow((1495,240),(1180,390))
d.text((70,950),'Стрелки показывают зависимости исходного кода. Статус элементов: реализовано в рабочей ветке.',font=map_font,fill='#475569')
img.save(figdir/'figure4.png')

# Evidence image based on the captured test runner summary; not a fabricated IDE screenshot.
code_card(figdir/'figure5.png', 'Фрагмент фактического вывода dotnet test', [
    ('Команда', 'dotnet test tests\\QuizMvp.Tests\\QuizMvp.Tests.csproj --no-restore --verbosity quiet'),
    ('Результат запуска', '''Пройдено: 25
Не пройдено: 0
Пропущено: 0
Всего: 25
Длительность: около 1 секунды
Сборка прошла с предупреждениями существующего кода.'''),
    ('Набор подмены', '''Вход корректный: рабочая + тестовая реализация → 302, cookie.
Пароль неверный: обе → форма входа, без cookie.
Имя отсутствует: обе → форма входа, без cookie.''')
])
code_card(figdir/'figure6.png', 'ПРОТОКОЛ ВЗАИМНОЙ ПРОВЕРКИ', [
    ('Текущий статус', '''НЕ ВЫПОЛНЕНО — проверяющий студент не назначен.
Версия чужого модуля: не предоставлена.
Команды и фактические результаты: отсутствуют.
Снимок проверки: отсутствует.'''),
    ('Условие завершения раздела Р7', '''Получить назначение и версию чужого модуля.
Собрать проект по его инструкции.
Проверить контракт, подмену и модульную карту.
Заполнить протокол и приложить реальный снимок.''')
])

# Start from the supplied report template and fill its existing slots.
doc = Document(str(template))
paras = doc.paragraphs
paras[32].text = ('В этой работе выделен модуль регистрации и входа Quizzy. Контракт отделяет правила учётной записи и работу с хранилищем от HTTP-контроллера. Рабочая реализация и тестовая подстановка подключаются за одним интерфейсом. Интеграционные проверки прошли; взаимная проверка другого студента ожидает назначения.')
paras[37].text = ('Цель работы: выделить контракт регистрации и входа, реализовать его в проекте Quizzy и проверить, что один и тот же потребитель работает с рабочей реализацией и тестовой подстановкой.')
paras[38].text = ('Ответственность модуля (Р1): FR-01 регистрация → создание учётной записи с проверкой занятого имени; FR-02 вход → проверка имени и пароля с типизированным исходом; публичные операции RegisterAsync и AuthenticateAsync → потребитель AccountController. HTTP-формы, перенаправления и пользовательская сессия остаются вне модуля.')
paras[40].text = ('Проект Quizzy — ASP.NET Core MVC приложение на .NET 8 с SQLite и EF Core. Модуль расположен в QuizMvp/Authentication и QuizMvp/Services. AccountController зависит только от IAuthenticationModule; конкретная реализация выбирается в Program.cs. Рабочая ветка: feature/quizmvp-modular-architecture; изменения локальные, не отправлялись и не объединялись.')
paras[44].text = ('Контракт определён до реализации и хранится рядом с исходным кодом: XML-комментарии находятся у публичного API, подробное описание — в docs/authentication-module-contract.md. Он задаёт вход, типизированные результаты, порядок проверок, передачу отмены, владение данными и побочные эффекты.')
paras[45].text = ('Ключевые границы: при входе пустой пароль и отсутствие пользователя дают InvalidCredentials; при регистрации сначала проверяется настройка, затем занятость имени. Отменённый запрос или сбой БД не маскируется и передаётся вызывающему коду. Модуль не изменяет HTTP-сессию.')
paras[49].text = ('Рабочий AuthenticationModule читает и записывает пользователей через ApplicationDbContext. Контроллер преобразует результаты контракта в представления, код HTTP и сессию. DI-регистрация добавлена в Program.cs; диагностический маршрут ПЗ1 сохранён. Существующее хеширование SHA-256 оставлено ради совместимости и отмечено как ограничение MVP.')
paras[54].text = ('Для замены используется один потребитель AccountController и один HTTP-сценарий входа. В production DI подключает AuthenticationModule. В тестовой точке сборки сервис удаляется и подключается SuccessfulLoginSubstitute. Формы и код контроллера одинаковы для обеих реализаций.')
paras[55].text = ('Подстановка возвращает заранее заданный аккаунт только для тестовой пары имени и пароля. Она проверяет обработку результата контроллером, но не проверяет SQLite, проверку хеша или сохранение учётной записи. Для этих свойств отдельно выполняются интеграционные тесты с рабочим модулем.')
paras[59].text = ('На карте зафиксированы зависимости, проверенные по using, вызовам и DI-регистрации: AccountController → IAuthenticationModule; AuthenticationModule реализует контракт и зависит от настроек и ApplicationDbContext; контекст использует модель User. Циклов для выделенного контура не обнаружено. Полная карта с расположением и статусами внесена в docs/MODULE_MAP.md.')
paras[64].text = ('Направление 1 — межкомпонентное взаимодействие: AccountController → контракт → рабочая или тестовая реализация. Вопрос: меняется ли поведение HTTP-потребителя при переключении реализации за контрактом? Готовность высокая для локального сценария, но тестовая подстановка не моделирует внешнюю инфраструктуру.')
paras[65].text = ('Направление 2 — хранение и обработка данных: поиск пользователя, обработка хеша пароля и чтение/запись SQLite. Вопрос: как объём учётных записей влияет на время поиска по имени? Готовность базовая: есть SQLite и наблюдаемые результаты, замеры производительности не проводились.')
paras[66].text = ('Модуль-кандидат — AuthenticateAsync. Воспроизводимый вход: существующий тестовый пользователь и правильный/неправильный пароль; наблюдаемые исходы — Authenticated/InvalidCredentials и HTTP-ответ. ADR-002 (автор Егор Голубев) сравнивает сохранение логики в контроллере, контракт с реализацией и перенос HTTP-логики в модуль; выбран контракт. Цена — дополнительные типы и косвенный вызов. Проверка — один HTTP-сценарий на обеих реализациях.')
paras[70].text = ('Проверки выполнены командой dotnet test tests/QuizMvp.Tests/QuizMvp.Tests.csproj --no-restore --verbosity quiet. Все 25 тестов прошли, 0 ошибок и 0 пропусков. В том числе проверены вход и границы контракта на рабочей и тестовой реализациях, конфигурация регистрации, сохранение пользователя и отказ при неверном пароле.')
paras[74].text = ('Р7 не выполнен: проверяющий студент и версия его модуля пока не назначены. При индивидуальном выполнении методичка требует модуль студента другого проекта по назначению преподавателя. Ниже зафиксирован статус ожидания; протокол проверки и снимок нельзя заполнить до получения назначения и фактического запуска.')
paras[78].text = ('В проекте выделен и интегрирован модуль регистрации и входа с документированным контрактом. AccountController зависит от контракта и принимает рабочую или тестовую реализацию без изменения сценария. Все 25 локальных тестов прошли. Проверка другого студента остаётся обязательным незавершённым пунктом; до неё отчет не является полностью готовым к сдаче.')

# Figures go into the matching template figure slots.
for index, filename in [(46,'figure2.png'),(50,'figure1.png'),(56,'figure3.png'),(60,'figure4.png'),(71,'figure5.png'),(75,'figure6.png')]:
    p = paras[index]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(figdir/filename), width=Inches(6.1))
paras[76].text = 'Рисунок 6 — Статус взаимной проверки (проверка ожидает назначения)'

# Fill the existing template tables without changing their layout.
def set_cell(table, row, col, value):
    cell = table.cell(row,col)
    cell.text = value
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

meta = doc.tables[3]
for row, value in enumerate([
    'Егор Голубев; группа не указана',
    'Quizzy (QuizMvp), .NET 8 / ASP.NET Core MVC / SQLite',
    'https://github.com/yunnanke/QuizVK (рабочая локальная копия)',
    'Регистрация и аутентификация',
    'Кандидат 0104-pz02-golubev-v1; тег и итоговый commit не созданы'
], start=1): set_cell(meta,row,1,value)

contract = doc.tables[4]
contract_values = [
    'RegisterAsync(request, token); AuthenticateAsync(request, token); RegistrationEnabled.',
    'Запросы содержат имя, email, пароль, роль; поля формы могут быть null. CancellationToken передаётся операции.',
    'Регистрация: Registered + аккаунт, Disabled или UsernameAlreadyExists. Вход: Authenticated + аккаунт либо InvalidCredentials.',
    'Регистрация: настройка → занятое имя → сохранение. Вход: поиск имени → проверка наличия пароля → хеш. Сбой хранилища и отмена передаются вызывающему коду.',
    'Входные записи не изменяются. Модуль владеет контекстом БД; наружу отдаёт ID, имя и роль, нужные для сессии.',
    'Регистрация записывает одну запись; вход только читает. HTTP-сессией управляет AccountController.'
]
for row,value in enumerate(contract_values,start=1): set_cell(contract,row,1,value)

subs = doc.tables[5]
sub_values = [
    ('Authentication\nModule', 'Рабочий модуль: SQLite/EF Core, регистрация, проверка пароля и вход.', 'Зависит от реальной БД; SHA-256 сохранён как текущее ограничение MVP.'),
    ('SuccessfulLogin\nSubstitute (тесты)', 'Тестовая подстановка контракта для того же HTTP-сценария.', 'Не проверяет SQLite, хеширование, реальную регистрацию или устойчивость данных.')
]
for row, vals in enumerate(sub_values,start=1):
    for col,value in enumerate(vals): set_cell(subs,row,col,value)

complexity = doc.tables[6]
complexity_values = [
    ('Межкомпонентное взаимодействие', 'AccountController → IAuthentication\nModule → DI-реализация', 'Меняется ли наблюдаемое поведение потребителя при переключении реализации?', 'Высокая для локального входа; внешние отказы не моделируются.'),
    ('Хранение и обработка данных', 'Поиск пользователя и проверка хеша через SQLite/EF Core', 'Как объём пользователей влияет на время AuthenticateAsync?', 'Базовая; реализация есть, замеры скорости не выполнялись.')
]
for row, vals in enumerate(complexity_values,start=1):
    for col,value in enumerate(vals): set_cell(complexity,row,col,value)

results = doc.tables[7]
result_values = [
    ('Правильные имя и пароль, рабочий модуль', '302 на главную, cookie сессии', 'Пройдено'),
    ('Те же данные, тестовая подстановка', 'Тот же 302 и cookie', 'Пройдено'),
    ('Неверный пароль, рабочий модуль', 'Форма входа, без cookie', 'Пройдено'),
    ('Неверный пароль, тестовая подстановка', 'Форма входа, без cookie', 'Пройдено'),
    ('Имя отсутствует, обе реализации', 'Форма входа, без cookie', 'Пройдено для обеих'),
    ('Регистрация выключена настройкой', 'GET/POST возвращают 503', 'Пройдено'),
    ('Повторная регистрация того же имени', 'Вторая запись не создаётся', 'Пройдено')
]
for row_index, vals in enumerate(result_values):
    if row_index + 1 < len(results.rows):
        target_row = row_index + 1
    else:
        results.add_row()
        target_row = len(results.rows) - 1
    for col,value in enumerate(vals): set_cell(results,target_row,col,value)

# Replace the four numbered source slots from the supplied template.
refs = [
    'Методические указания к практической работе № 2 «Модульная архитектура и контракт программного модуля», МДК.01.04, предоставленная версия.',
    'Шаблон отчёта к практической работе № 2, предоставленная версия.',
    'Quizzy: docs/REQUIREMENTS.md, docs/authentication-module-contract.md, docs/MODULE_MAP.md и docs/adr/ADR-002-authentication-module.md; локальная рабочая версия.',
    'Исходный код QuizMvp и проверки tests/QuizMvp.Tests; команда и фактический результат приведены в разделе 4.1.'
]
for i, value in enumerate(refs, start=80): paras[i].text = f'{i-79}) {value}'

# Keep the supplied style hierarchy and set deferred field refresh (TOC/page references).
settings = doc.settings.element
update = settings.find(qn('w:updateFields'))
if update is None:
    update = OxmlElement('w:updateFields')
    settings.append(update)
update.set(qn('w:val'),'true')

# Fill the student's name and normalize table paragraphs for readable wrapping.
doc.tables[2].cell(0,5).text = 'Е. Е. Голубев'
for table in doc.tables[3:]:
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.size = Pt(10)

out.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(out))
print(out)


